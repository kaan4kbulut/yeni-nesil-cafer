"""inspect_output aracının yardımcıları: bir çıktı dosyasını görme modelinin bakabileceği resme çevirir.

Model "yaptım" dediğinde sonucun istenen şey olup olmadığını gözle denetlesin diye (agent._tool_inspect_output bu
resmi look_at_image'a verir). Resme çevrilemeyen belgeler (Word, PowerPoint, Excel) için yapı raporu döner.
3D ve video çizimi ajanların Python'unda (trimesh, matplotlib, opencv) ayrı süreçte yapılır: programın süreci ağır
kütüphaneleri yüklemez.
"""

import subprocess
from pathlib import Path

IMAGE = {".png", ".jpg", ".jpeg", ".webp", ".bmp", ".gif"}
MODEL = {".stl", ".obj", ".glb", ".gltf", ".ply", ".3mf", ".off"}
VIDEO = {".mp4", ".webm", ".mkv", ".mov", ".avi", ".m4v"}

# 3D modeli üç açıdan çizer (perspektif, önden, üstten); ölçüleri başlığa yazar
_RENDER_3D = r"""
import sys
import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
import numpy as np
import trimesh
from mpl_toolkits.mplot3d.art3d import Poly3DCollection
mesh = trimesh.load(sys.argv[1], force="mesh")
if len(mesh.faces) > 60000:
    mesh = mesh.simplify_quadric_decimation(face_count=60000) if hasattr(mesh, "simplify_quadric_decimation") else mesh
tri = mesh.vertices[mesh.faces]
light = np.array([0.4, -0.5, 0.8]); light /= np.linalg.norm(light)
shade = 0.35 + 0.65 * np.clip(mesh.face_normals @ light, 0, 1)
colors = np.c_[0.25 * shade, 0.55 * shade, 0.9 * shade, np.ones_like(shade)]
lo, hi = mesh.bounds; mid = (lo + hi) / 2; r = (hi - lo).max() / 2
x, y, z = (hi - lo).round(2)
fig = plt.figure(figsize=(12, 4.4), dpi=110)
for i, (title, elev, azim) in enumerate([("perspektif", 25, -55), ("önden", 0, -90), ("üstten", 90, -90)]):
    ax = fig.add_subplot(1, 3, i + 1, projection="3d")
    ax.add_collection3d(Poly3DCollection(tri, facecolors=colors, edgecolors="none"))
    for set_lim, m in ((ax.set_xlim, mid[0]), (ax.set_ylim, mid[1]), (ax.set_zlim, mid[2])):
        set_lim(m - r, m + r)
    ax.set_box_aspect((1, 1, 1)); ax.view_init(elev=elev, azim=azim); ax.set_title(title); ax.set_axis_off()
fig.suptitle(f"{x} x {y} x {z} mm")
fig.tight_layout(); fig.savefig(sys.argv[2]); print("ok")
"""

# videonun ortasından bir kare
_VIDEO_FRAME = r"""
import sys, cv2
cap = cv2.VideoCapture(sys.argv[1])
n = int(cap.get(cv2.CAP_PROP_FRAME_COUNT) or 0); fps = cap.get(cv2.CAP_PROP_FPS) or 0
cap.set(cv2.CAP_PROP_POS_FRAMES, max(0, n // 2))
ok, frame = cap.read()
if not ok: raise SystemExit("video okunamadı")
cv2.imwrite(sys.argv[2], frame)
h, w = frame.shape[:2]
print(f"{w}x{h}, {n} kare, {fps:.0f} fps, {n / fps if fps else 0:.1f} sn")
"""


def _run(code: str, *args: str) -> tuple[bool, str]:
    from .tools import NO_WINDOW, agent_env, python_exe

    out = subprocess.run([python_exe(), "-c", code, *args], capture_output=True, text=True, encoding="utf-8", errors="replace", timeout=180,
                         env=agent_env(), creationflags=NO_WINDOW)
    if out.returncode != 0:
        return False, (out.stderr.strip().splitlines() or ["?"])[-1]
    return True, out.stdout.strip()


def _pdf(path: Path, png: Path) -> str:
    """PDF'in ilk iki sayfası yan yana (Qt'nin PDF motoru; ek program gerekmez)."""
    from PySide6.QtCore import QSize
    from PySide6.QtGui import QImage, QPainter
    from PySide6.QtPdf import QPdfDocument

    doc = QPdfDocument()
    if doc.load(str(path)) != QPdfDocument.Error.None_:
        raise ValueError("PDF açılamadı")
    pages = [doc.render(i, QSize(700, int(700 * doc.pagePointSize(i).height() / doc.pagePointSize(i).width())))
             for i in range(min(2, doc.pageCount()))]
    canvas = QImage(sum(p.width() for p in pages) + 10 * (len(pages) - 1), max(p.height() for p in pages),
                    QImage.Format_RGB32)
    canvas.fill(0xFFFFFF)
    painter = QPainter(canvas)
    x = 0
    for p in pages:
        painter.drawImage(x, 0, p)
        x += p.width() + 10
    painter.end()
    canvas.save(str(png))
    return f"{doc.pageCount()} pages (first {len(pages)} shown)"


def _svg(path: Path, png: Path) -> str:
    from PySide6.QtCore import Qt
    from PySide6.QtGui import QImage, QPainter
    from PySide6.QtSvg import QSvgRenderer

    r = QSvgRenderer(str(path))
    size = r.defaultSize().scaled(1000, 1000, Qt.KeepAspectRatio)
    img = QImage(size, QImage.Format_ARGB32)
    img.fill(0xFFFFFFFF)
    painter = QPainter(img)
    r.render(painter)
    painter.end()
    img.save(str(png))
    return f"SVG {size.width()}x{size.height()}"


def _structure(path: Path) -> str:
    """Resme çevrilemeyen ofis belgesinin yapısı (ajanların Python'unda)."""
    code = r"""
import sys, json
p = sys.argv[1]; kind = p.rsplit(".", 1)[-1].lower(); r = {}
if kind == "docx":
    from docx import Document
    d = Document(p)
    r = {"headings": [x.text for x in d.paragraphs if x.style.name.lower().startswith(("heading", "title", "başlık"))][:20],
         "paragraphs": sum(1 for x in d.paragraphs if x.text.strip()), "tables": len(d.tables),
         "pictures": len(d.inline_shapes), "first_text": " ".join(x.text for x in d.paragraphs if x.text.strip())[:600]}
elif kind == "pptx":
    from pptx import Presentation
    s = Presentation(p).slides
    r = {"slides": [{"title": (sl.shapes.title.text if sl.shapes.title is not None else ""),
                     "pictures": sum(1 for sh in sl.shapes if sh.shape_type == 13),
                     "text": " ".join(sh.text_frame.text for sh in sl.shapes if sh.has_text_frame)[:200]} for sl in s][:20]}
elif kind in ("xlsx", "xlsm"):
    from openpyxl import load_workbook
    wb = load_workbook(p)
    r = {"sheets": [{"name": ws.title, "size": ws.dimensions, "images": len(ws._images), "charts": len(ws._charts),
                     "first_rows": [[c for c in row] for row in ws.iter_rows(max_row=4, values_only=True)]}
                    for ws in wb.worksheets]}
print(json.dumps(r, ensure_ascii=False, default=str))
"""
    ok, out = _run(code, str(path))
    if not ok:
        raise ValueError(out)
    return out


def render(path: Path, out_dir: Path) -> tuple[Path | None, str]:
    """(resim ya da None, bilgi metni). Resim görme modeline gider; bilgi metni modele olduğu gibi."""
    suffix = path.suffix.lower()
    out_dir.mkdir(parents=True, exist_ok=True)
    png = out_dir / f"{path.stem}-denetim.png"
    if suffix in IMAGE:
        return path, ""
    if suffix == ".svg":
        return png, _svg(path, png)
    if suffix == ".pdf":
        return png, _pdf(path, png)
    if suffix in MODEL:
        ok, out = _run(_RENDER_3D, str(path), str(png))
        if not ok:
            raise ValueError(f"3D model could not be drawn: {out}")
        return png, "3 views (perspective, front, top); size in the title"
    if suffix in VIDEO:
        ok, out = _run(_VIDEO_FRAME, str(path), str(png))
        if not ok:
            raise ValueError(f"video could not be read: {out}")
        return png, f"middle frame; {out}"
    if suffix in (".docx", ".pptx", ".xlsx", ".xlsm"):
        facts = _structure(path)
        return None, "Structure (office files cannot be rendered to an image here): " + facts
    raise ValueError(f"cannot inspect {suffix or 'this'} files; look at them with read_file instead")

